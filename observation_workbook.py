from __future__ import annotations

from copy import copy
from datetime import date as date_type, datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook


# ============================================================
# Configuration
# ============================================================

WORKBOOK_DIR = Path.home() / "Documents" / "AI_investing_observation"

SOURCE_FILE = WORKBOOK_DIR / "AI_investing_observation.xlsx"

TEST_FILE = WORKBOOK_DIR / "AI_investing_observation_test.xlsx"

DAILY_RUN_SHEET = "Daily_Run"

DAILY_RUN_HEADER_ROW = 4
DAILY_RUN_FIRST_DATA_ROW = 5

DAILY_RUN_HEADERS = [
    "Date",
    "Version",
    "PipelineStatus",
    "PassSteps",
    "RunId",
    "AsOfDate",
    "UniverseVersion",
    "ScoreModelVersion",
    "RiskModelVersion",
    "UniverseConfigured",
    "Ready",
    "Excluded",
    "ProviderRejected",
    "StaleMarketData",
    "InsufficientHistory",
    "ExcludedSymbols",
    "CoverageStatus",
    "BUY",
    "WATCH",
    "IGNORE",
    "FinalStatus",
    "Notes",
]


# ============================================================
# Validation
# ============================================================

def validate_daily_run_schema(ws) -> None:
    """
    Confirm that Daily_Run still has the expected schema.

    Fail closed if the workbook structure has changed.
    """

    actual_headers = [
        ws.cell(row=DAILY_RUN_HEADER_ROW, column=col).value
        for col in range(1, len(DAILY_RUN_HEADERS) + 1)
    ]

    if actual_headers != DAILY_RUN_HEADERS:
        print("\nERROR: Daily_Run schema mismatch.")
        print("\nExpected:")
        print(DAILY_RUN_HEADERS)

        print("\nActual:")
        print(actual_headers)

        raise RuntimeError(
            "Daily_Run schema validation failed. Workbook was NOT modified."
        )


def _normalize_daily_run_date(value):
    """Normalize Excel/Python date values for duplicate-date checks."""

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date_type):
        return value

    if isinstance(value, str):
        try:
            return date_type.fromisoformat(value.strip())
        except ValueError:
            return value.strip()

    return value


def validate_daily_run_history_contiguous(ws) -> None:
    """
    Confirm that Daily_Run history is append-only with no internal blank row.

    A blank row followed by a later populated row indicates a history gap.
    Fail closed rather than backfilling an older position.
    """

    first_empty_row = None

    for row in range(DAILY_RUN_FIRST_DATA_ROW, ws.max_row + 1):
        values = [
            ws.cell(row=row, column=col).value
            for col in range(1, len(DAILY_RUN_HEADERS) + 1)
        ]

        row_is_empty = all(value is None for value in values)

        if row_is_empty:
            if first_empty_row is None:
                first_empty_row = row
            continue

        if first_empty_row is not None:
            raise RuntimeError(
                "Daily_Run history gap detected: "
                f"row {first_empty_row} is empty but row {row} contains data. "
                "Workbook was NOT modified."
            )


def validate_daily_run_duplicate(ws, date, run_id) -> None:
    """
    Reject a duplicate observation Date or RunId.

    Daily_Run policy is one record per day and RunId must be unique.
    """

    target_date = _normalize_daily_run_date(date)

    for row in range(DAILY_RUN_FIRST_DATA_ROW, ws.max_row + 1):
        existing_date = ws.cell(row=row, column=1).value
        existing_run_id = ws.cell(row=row, column=5).value

        if existing_date is None and existing_run_id is None:
            continue

        if _normalize_daily_run_date(existing_date) == target_date:
            raise RuntimeError(
                f"Duplicate Daily_Run Date detected: {target_date}. "
                f"Existing row: {row}. Workbook was NOT modified."
            )

        if existing_run_id == run_id:
            raise RuntimeError(
                f"Duplicate Daily_Run RunId detected: {run_id}. "
                f"Existing row: {row}. Workbook was NOT modified."
            )


def validate_daily_run_values(
    *,
    universe_configured,
    ready,
    excluded,
    buy,
    watch,
    ignore,
) -> None:
    """Validate internal Daily_Run arithmetic consistency."""

    if ready + excluded != universe_configured:
        raise RuntimeError(
            "Daily_Run value validation failed: "
            f"Ready({ready}) + Excluded({excluded}) "
            f"!= UniverseConfigured({universe_configured}). "
            "Workbook was NOT modified."
        )

    if buy + watch + ignore != ready:
        raise RuntimeError(
            "Daily_Run value validation failed: "
            f"BUY({buy}) + WATCH({watch}) + IGNORE({ignore}) "
            f"!= Ready({ready}). "
            "Workbook was NOT modified."
        )


# ============================================================
# Row discovery
# ============================================================

def find_first_empty_daily_run_row(ws) -> int:
    """
    Find the first truly empty Daily_Run data row.

    Do NOT rely on ws.max_row because the workbook contains
    preformatted / reserved rows.
    """

    for row in range(DAILY_RUN_FIRST_DATA_ROW, ws.max_row + 2):

        values = [
            ws.cell(row=row, column=col).value
            for col in range(1, len(DAILY_RUN_HEADERS) + 1)
        ]

        if all(value is None for value in values):
            return row

    raise RuntimeError("No empty Daily_Run row found.")


# ============================================================
# Formatting
# ============================================================

def copy_row_format(ws, source_row: int, target_row: int) -> None:
    """
    Copy formatting only.
    Do NOT copy historical values.
    """

    for col in range(1, len(DAILY_RUN_HEADERS) + 1):

        source_cell = ws.cell(row=source_row, column=col)
        target_cell = ws.cell(row=target_row, column=col)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format

        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)

    if ws.row_dimensions[source_row].height:
        ws.row_dimensions[target_row].height = (
            ws.row_dimensions[source_row].height
        )


# ============================================================
# Backup
# ============================================================

def make_backup(file_path: Path) -> Path:

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")

    backup_path = file_path.with_name(
        f"{file_path.stem}.backup-{timestamp}{file_path.suffix}"
    )

    shutil.copy2(file_path, backup_path)

    print(f"Backup created: {backup_path}")

    return backup_path


# ============================================================
# Daily_Run writer
# ============================================================

def append_daily_run(
    file_path: Path,
    *,
    date,
    version,
    pipeline_status,
    pass_steps,
    run_id,
    as_of_date,
    universe_version,
    score_model_version,
    risk_model_version,
    universe_configured,
    ready,
    excluded,
    provider_rejected,
    stale_market_data,
    insufficient_history,
    excluded_symbols,
    coverage_status,
    buy,
    watch,
    ignore,
    final_status,
    notes,
) -> int:

    if not file_path.exists():
        raise FileNotFoundError(file_path)

    # --------------------------------------------------------
    # Load
    # --------------------------------------------------------

    wb = load_workbook(file_path)

    if DAILY_RUN_SHEET not in wb.sheetnames:
        raise RuntimeError(
            f"Required sheet '{DAILY_RUN_SHEET}' does not exist."
        )

    ws = wb[DAILY_RUN_SHEET]

    # --------------------------------------------------------
    # Schema validation
    # --------------------------------------------------------

    validate_daily_run_schema(ws)
    validate_daily_run_history_contiguous(ws)
    validate_daily_run_duplicate(
        ws,
        date=date,
        run_id=run_id,
    )
    validate_daily_run_values(
        universe_configured=universe_configured,
        ready=ready,
        excluded=excluded,
        buy=buy,
        watch=watch,
        ignore=ignore,
    )

    # --------------------------------------------------------
    # Find safe append row
    # --------------------------------------------------------

    target_row = find_first_empty_daily_run_row(ws)

    # Preserve formatting using previous populated row
    previous_row = target_row - 1

    if previous_row >= DAILY_RUN_FIRST_DATA_ROW:
        copy_row_format(ws, previous_row, target_row)

    # --------------------------------------------------------
    # Prepare values
    # --------------------------------------------------------

    values = [
        date,
        version,
        pipeline_status,
        pass_steps,
        run_id,
        as_of_date,
        universe_version,
        score_model_version,
        risk_model_version,
        universe_configured,
        ready,
        excluded,
        provider_rejected,
        stale_market_data,
        insufficient_history,
        excluded_symbols,
        coverage_status,
        buy,
        watch,
        ignore,
        final_status,
        notes,
    ]

    # --------------------------------------------------------
    # Write
    # --------------------------------------------------------

    for col, value in enumerate(values, start=1):
        ws.cell(row=target_row, column=col).value = value

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    wb.save(file_path)

    # --------------------------------------------------------
    # Verification: reopen
    # --------------------------------------------------------

    verify_wb = load_workbook(file_path, data_only=False)
    verify_ws = verify_wb[DAILY_RUN_SHEET]

    verify_values = [
        verify_ws.cell(row=target_row, column=col).value
        for col in range(1, len(DAILY_RUN_HEADERS) + 1)
    ]

    if verify_values != values:
        raise RuntimeError(
            "Post-save verification failed."
        )

    print("\nDaily_Run append PASS")
    print(f"Workbook : {file_path}")
    print(f"Sheet    : {DAILY_RUN_SHEET}")
    print(f"Row      : {target_row}")
    print(f"Date     : {date}")
    print(f"RunId    : {run_id}")
    print(f"READY    : {ready}/{universe_configured}")
    print(f"Status   : {final_status}")

    return target_row


# ============================================================
# Test copy
# ============================================================

def prepare_test_workbook() -> Path:

    if not SOURCE_FILE.exists():
        raise FileNotFoundError(
            f"Original workbook not found:\n{SOURCE_FILE}"
        )

    shutil.copy2(SOURCE_FILE, TEST_FILE)

    print(f"Test workbook created:\n{TEST_FILE}")

    return TEST_FILE


# ============================================================
# Main test
# ============================================================

if __name__ == "__main__":

    test_file = prepare_test_workbook()

    make_backup(test_file)

    append_daily_run(
        test_file,

        # TEST DATA ONLY
        date=datetime(2099, 1, 1),

        version="TEST",

        pipeline_status="PASS",

        pass_steps=20,

        run_id="TEST-RUN-ID",

        as_of_date=datetime(2098, 12, 31),

        universe_version="TEST-UNIVERSE",

        score_model_version="TEST-SCORE-MODEL",

        risk_model_version="MISSING",

        universe_configured=150,

        ready=148,

        excluded=2,

        provider_rejected=0,

        stale_market_data=0,

        insufficient_history=2,

        excluded_symbols="TEST1, TEST2",

        coverage_status="TEST_ONLY",

        buy=0,

        watch=1,

        ignore=147,

        final_status="NO_ACTION",

        notes="AUTOMATION TEST ROW — SAFE TO DELETE FROM TEST COPY.",
    )