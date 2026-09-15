"""Controlled historical remediation for Observation Workbook row 2026-09-15.

Scope is intentionally frozen:
- Daily_Run: RiskModelVersion, CoverageStatus, Notes only.
- Candidate_Tracking: FundamentalScore, CombinedScore only for OKTA and CRWD.
- Maintenance_Log: append one audit record dated at execution time.
- No other historical values may change.

Default/audit and --write-test never modify the production workbook.
Production write requires both --write-production and --confirm-production-write.
"""

from __future__ import annotations

import argparse
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from observation_runner import (
    build_daily_run_notes,
    classify_observation_coverage_status,
)
from observation_workbook import append_maintenance_log


OBSERVATION_DIR = Path.home() / "Documents" / "AI_investing_observation"
PRODUCTION_WORKBOOK = OBSERVATION_DIR / "AI_investing_observation.xlsx"
TEST_WORKBOOK = (
    OBSERVATION_DIR / "AI_investing_observation_historical_remediation_test.xlsx"
)
TEMP_WORKBOOK = (
    OBSERVATION_DIR / "AI_investing_observation.historical-remediation.tmp.xlsx"
)
BACKUP_DIR = OBSERVATION_DIR / "backups"

TARGET_DATE = date(2026, 9, 15)
TARGET_RUN_ID = "candidate-20260914-93aea4566066"
TARGET_TICKERS = ("OKTA", "CRWD")

REMEDIATION_COMMITS = {
    "RiskModelVersion": "821eff70",
    "Notes": "84d1f0b",
    "CoverageStatus": "30915ca",
    "MissingScores": "d757f31",
}

DAILY_ALLOWED_FIELDS = {
    "RiskModelVersion",
    "CoverageStatus",
    "Notes",
}
CANDIDATE_ALLOWED_FIELDS = {
    "FundamentalScore",
    "CombinedScore",
}


def _normalize_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return date.fromisoformat(value.strip())
    raise RuntimeError(f"Unsupported date value: {value!r}")


def _headers(ws, header_row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=col).value
        if value is None:
            continue
        key = str(value).strip()
        if key in result:
            raise RuntimeError(f"Duplicate header {key!r} in {ws.title}")
        result[key] = col
    return result


def _require_headers(
    ws,
    header_row: int,
    required: tuple[str, ...],
) -> dict[str, int]:
    mapping = _headers(ws, header_row)
    missing = [name for name in required if name not in mapping]
    if missing:
        raise RuntimeError(
            f"{ws.title} missing required headers: {', '.join(missing)}"
        )
    return mapping


def _find_daily_row(ws) -> tuple[int, dict[str, int]]:
    required = (
        "Date",
        "RunId",
        "FinalStatus",
        "UniverseConfigured",
        "Ready",
        "Excluded",
        "ProviderRejected",
        "StaleMarketData",
        "InsufficientHistory",
        "ExcludedSymbols",
        "RiskModelVersion",
        "CoverageStatus",
        "Notes",
    )
    cols = _require_headers(ws, 4, required)

    matches: list[int] = []
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row=row, column=cols["Date"]).value
        if value is None:
            continue
        try:
            row_date = _normalize_date(value)
        except Exception:
            continue
        if row_date == TARGET_DATE:
            matches.append(row)

    if len(matches) != 1:
        raise RuntimeError(
            f"Daily_Run expected exactly one {TARGET_DATE} row; found {matches}"
        )

    row = matches[0]
    run_id = str(ws.cell(row=row, column=cols["RunId"]).value or "").strip()
    if run_id != TARGET_RUN_ID:
        raise RuntimeError(
            f"Daily_Run RunId gate failed: {run_id!r} != {TARGET_RUN_ID!r}"
        )

    return row, cols


def _find_candidate_row(
    ws,
    ticker: str,
) -> tuple[int, dict[str, int]]:
    required = (
        "Date",
        "Ticker",
        "FundamentalScore",
        "CombinedScore",
        "WhyTracked",
        "ResearchNote",
    )
    cols = _require_headers(ws, 4, required)

    matches: list[int] = []
    for row in range(5, ws.max_row + 1):
        raw_date = ws.cell(row=row, column=cols["Date"]).value
        raw_ticker = ws.cell(row=row, column=cols["Ticker"]).value
        if raw_date is None or raw_ticker is None:
            continue
        try:
            row_date = _normalize_date(raw_date)
        except Exception:
            continue
        symbol = str(raw_ticker).strip().upper()
        if row_date == TARGET_DATE and symbol == ticker:
            matches.append(row)

    if len(matches) != 1:
        raise RuntimeError(
            f"Candidate_Tracking expected exactly one {TARGET_DATE}/{ticker} row; "
            f"found {matches}"
        )
    return matches[0], cols


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def compute_daily_updates(ws) -> tuple[int, dict[str, Any], dict[str, Any]]:
    row, cols = _find_daily_row(ws)

    final_status = str(
        ws.cell(row=row, column=cols["FinalStatus"]).value or ""
    ).strip().upper()
    if final_status != "NO_ACTION":
        raise RuntimeError(
            f"RiskModelVersion remediation requires NO_ACTION; got {final_status!r}"
        )

    configured = int(ws.cell(row=row, column=cols["UniverseConfigured"]).value)
    ready = int(ws.cell(row=row, column=cols["Ready"]).value)
    excluded = int(ws.cell(row=row, column=cols["Excluded"]).value)
    provider_rejected = int(
        ws.cell(row=row, column=cols["ProviderRejected"]).value
    )
    stale_market_data = int(
        ws.cell(row=row, column=cols["StaleMarketData"]).value
    )
    insufficient_history = int(
        ws.cell(row=row, column=cols["InsufficientHistory"]).value
    )
    excluded_symbols = str(
        ws.cell(row=row, column=cols["ExcludedSymbols"]).value or ""
    ).strip()

    before = {
        field: ws.cell(row=row, column=cols[field]).value
        for field in DAILY_ALLOWED_FIELDS
    }

    coverage_status = classify_observation_coverage_status(
        configured=configured,
        ready=ready,
        excluded=excluded,
        provider_rejected=provider_rejected,
        stale_market_data=stale_market_data,
        insufficient_history=insufficient_history,
        excluded_symbols=excluded_symbols,
    )
    notes = build_daily_run_notes(
        configured=configured,
        ready=ready,
        excluded=excluded,
        provider_rejected=provider_rejected,
        stale_market_data=stale_market_data,
        insufficient_history=insufficient_history,
        excluded_symbols=excluded_symbols,
    )

    after = {
        "RiskModelVersion": "MISSING",
        "CoverageStatus": coverage_status,
        "Notes": notes,
    }

    return row, before, after


def compute_candidate_updates(
    ws,
) -> dict[str, tuple[int, dict[str, Any], dict[str, Any]]]:
    result: dict[str, tuple[int, dict[str, Any], dict[str, Any]]] = {}

    for ticker in TARGET_TICKERS:
        row, cols = _find_candidate_row(ws, ticker)

        before = {
            field: ws.cell(row=row, column=cols[field]).value
            for field in CANDIDATE_ALLOWED_FIELDS
        }

        if not all(_is_blank(value) for value in before.values()):
            if set(before.values()) == {"MISSING"}:
                raise RuntimeError(
                    f"{ticker} score fields are already remediated."
                )
            raise RuntimeError(
                f"{ticker} score gate failed; expected blank values, got {before}"
            )

        after = {
            "FundamentalScore": "MISSING",
            "CombinedScore": "MISSING",
        }
        result[ticker] = (row, before, after)

    return result


def build_plan(workbook_path: Path) -> dict[str, Any]:
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)

    wb = load_workbook(workbook_path, data_only=False)
    try:
        for sheet_name in ("Daily_Run", "Candidate_Tracking", "Maintenance_Log"):
            if sheet_name not in wb.sheetnames:
                raise RuntimeError(f"Required sheet missing: {sheet_name}")

        daily_row, daily_before, daily_after = compute_daily_updates(
            wb["Daily_Run"]
        )
        candidate_updates = compute_candidate_updates(wb["Candidate_Tracking"])

        return {
            "daily_row": daily_row,
            "daily_before": daily_before,
            "daily_after": daily_after,
            "candidates": candidate_updates,
        }
    finally:
        wb.close()


def print_plan(plan: dict[str, Any]) -> None:
    print("HISTORICAL OBSERVATION REMEDIATION — AUDIT")
    print("=" * 78)
    print(f"Target Date : {TARGET_DATE}")
    print(f"Target RunId: {TARGET_RUN_ID}")
    print()
    print(f"Daily_Run row {plan['daily_row']}")
    for field in ("RiskModelVersion", "CoverageStatus", "Notes"):
        print(f"  {field}")
        print(f"    BEFORE: {plan['daily_before'][field]!r}")
        print(f"    AFTER : {plan['daily_after'][field]!r}")

    for ticker in TARGET_TICKERS:
        row, before, after = plan["candidates"][ticker]
        print()
        print(f"Candidate_Tracking {ticker} row {row}")
        for field in ("FundamentalScore", "CombinedScore"):
            print(f"  {field}: {before[field]!r} -> {after[field]!r}")

    print("=" * 78)
    print("NO WORKBOOK WRITE WAS PERFORMED")


def _snapshot_existing_values(wb) -> dict[str, dict[tuple[int, int], Any]]:
    snapshot: dict[str, dict[tuple[int, int], Any]] = {}
    for ws in wb.worksheets:
        sheet_values: dict[tuple[int, int], Any] = {}
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                sheet_values[(row, col)] = ws.cell(row=row, column=col).value
        snapshot[ws.title] = sheet_values
    return snapshot


def _maintenance_record(
    *,
    execution_time: datetime,
    plan: dict[str, Any],
) -> dict[str, Any]:
    daily_before = plan["daily_before"]
    daily_after = plan["daily_after"]

    before_text = (
        "Daily_Run: "
        f"RiskModelVersion={daily_before['RiskModelVersion']!r}; "
        f"CoverageStatus={daily_before['CoverageStatus']!r}; "
        f"Notes={daily_before['Notes']!r}. "
        "Candidate_Tracking OKTA/CRWD: FundamentalScore=<blank>; "
        "CombinedScore=<blank>."
    )
    after_text = (
        "Daily_Run: "
        f"RiskModelVersion={daily_after['RiskModelVersion']!r}; "
        f"CoverageStatus={daily_after['CoverageStatus']!r}; "
        f"Notes={daily_after['Notes']!r}. "
        "Candidate_Tracking OKTA/CRWD: FundamentalScore='MISSING'; "
        "CombinedScore='MISSING'."
    )

    evidence = (
        "Semantic remediation commits: "
        f"RiskModelVersion={REMEDIATION_COMMITS['RiskModelVersion']}; "
        f"Notes={REMEDIATION_COMMITS['Notes']}; "
        f"CoverageStatus={REMEDIATION_COMMITS['CoverageStatus']}; "
        f"MissingScores={REMEDIATION_COMMITS['MissingScores']}. "
        f"Historical target Date={TARGET_DATE}, RunId={TARGET_RUN_ID}."
    )

    return {
        "date": datetime(
            execution_time.year, execution_time.month, execution_time.day
        ),
        "category": "CODE_BUG",
        "severity": "P1",
        "ticker_scope": "OBSERVATION 2026-09-15 / OKTA / CRWD",
        "issue": "Controlled correction of 2026-09-15 Observation automation semantics",
        "evidence": evidence,
        "action": "CODE_FIX",
        "file_changed": str(PRODUCTION_WORKBOOK.name),
        "before": before_text,
        "after": after_text,
        "result": "RESOLVED",
        "follow_up_date": None,
        "status": "CLOSED",
        "notes": (
            f"ExecutedAt={execution_time.isoformat(timespec='seconds')}; "
            f"HistoricalTargetDate={TARGET_DATE}; "
            "Only explicitly allowed cells were changed; WhyTracked, "
            "ResearchNote, horizons, Outcome and all other historical rows "
            "were protected."
        ),
    }


def _apply_to_copy(
    source: Path,
    destination: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    if destination.exists():
        destination.unlink()
    shutil.copy2(source, destination)

    plan = build_plan(destination)
    execution_time = datetime.now().astimezone()

    wb = load_workbook(destination, data_only=False)
    try:
        before_snapshot = _snapshot_existing_values(wb)

        daily_ws = wb["Daily_Run"]
        daily_cols = _headers(daily_ws, 4)
        daily_row = plan["daily_row"]
        for field, value in plan["daily_after"].items():
            daily_ws.cell(
                row=daily_row,
                column=daily_cols[field],
            ).value = value

        candidate_ws = wb["Candidate_Tracking"]
        candidate_cols = _headers(candidate_ws, 4)
        for ticker in TARGET_TICKERS:
            row, _, after = plan["candidates"][ticker]
            for field, value in after.items():
                candidate_ws.cell(
                    row=row,
                    column=candidate_cols[field],
                ).value = value

        wb.save(destination)
    finally:
        wb.close()

    maintenance = _maintenance_record(
        execution_time=execution_time,
        plan=plan,
    )
    maintenance_row = append_maintenance_log(destination, **maintenance)

    verify_wb = load_workbook(destination, data_only=False)
    try:
        # Verify targeted Daily_Run values.
        daily_ws = verify_wb["Daily_Run"]
        daily_cols = _headers(daily_ws, 4)
        for field, expected in plan["daily_after"].items():
            actual = daily_ws.cell(
                row=plan["daily_row"],
                column=daily_cols[field],
            ).value
            if actual != expected:
                raise RuntimeError(
                    f"Daily_Run {field} post-save verification failed: "
                    f"{actual!r} != {expected!r}"
                )

        # Verify targeted Candidate_Tracking values.
        candidate_ws = verify_wb["Candidate_Tracking"]
        candidate_cols = _headers(candidate_ws, 4)
        for ticker in TARGET_TICKERS:
            row, _, after = plan["candidates"][ticker]
            for field, expected in after.items():
                actual = candidate_ws.cell(
                    row=row,
                    column=candidate_cols[field],
                ).value
                if actual != expected:
                    raise RuntimeError(
                        f"{ticker} {field} post-save verification failed: "
                        f"{actual!r} != {expected!r}"
                    )

        # Verify every pre-existing cell value other than the explicitly
        # allowed target cells remained unchanged. Maintenance_Log is allowed
        # to gain exactly one append row; existing Maintenance_Log values must
        # still be unchanged.
        allowed: set[tuple[str, int, int]] = set()
        for field in DAILY_ALLOWED_FIELDS:
            allowed.add(
                (
                    "Daily_Run",
                    plan["daily_row"],
                    daily_cols[field],
                )
            )
        for ticker in TARGET_TICKERS:
            row, _, _ = plan["candidates"][ticker]
            for field in CANDIDATE_ALLOWED_FIELDS:
                allowed.add(
                    (
                        "Candidate_Tracking",
                        row,
                        candidate_cols[field],
                    )
                )

        for sheet_name, cells in before_snapshot.items():
            ws = verify_wb[sheet_name]
            for (row, col), before_value in cells.items():
                key = (sheet_name, row, col)
                if key in allowed:
                    continue
                # Maintenance_Log is append-only. The chosen append row may
                # already exist as a preformatted-but-empty row in Excel, so
                # its previously blank cells are intentionally allowed to
                # become the new audit record. All other existing
                # Maintenance_Log rows remain protected.
                if sheet_name == "Maintenance_Log" and row == maintenance_row:
                    continue
                after_value = ws.cell(row=row, column=col).value
                if after_value != before_value:
                    raise RuntimeError(
                        "Protected historical value changed unexpectedly: "
                        f"{sheet_name}!R{row}C{col}: "
                        f"{before_value!r} -> {after_value!r}"
                    )

        # Verify one Maintenance_Log record matching this remediation exists.
        maintenance_ws = verify_wb["Maintenance_Log"]
        maintenance_cols = _headers(maintenance_ws, 4)
        matches = []
        for row in range(5, maintenance_ws.max_row + 1):
            issue = maintenance_ws.cell(
                row=row,
                column=maintenance_cols["Issue"],
            ).value
            scope = maintenance_ws.cell(
                row=row,
                column=maintenance_cols["Ticker/Scope"],
            ).value
            if (
                issue == maintenance["issue"]
                and scope == maintenance["ticker_scope"]
            ):
                matches.append(row)

        if len(matches) != 1:
            raise RuntimeError(
                f"Expected one Maintenance_Log remediation record; found {matches}"
            )

        if matches[0] != maintenance_row:
            raise RuntimeError(
                "Maintenance_Log append row verification mismatch: "
                f"{matches[0]} != {maintenance_row}"
            )

        return plan, {
            "maintenance_row": maintenance_row,
            "execution_time": execution_time,
        }
    finally:
        verify_wb.close()


def _production_backup_path() -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    path = (
        BACKUP_DIR
        / f"AI_investing_observation.before-20260915-historical-remediation-{stamp}.xlsx"
    )
    if path.exists():
        raise RuntimeError(f"Backup path already exists: {path}")
    return path


def write_production() -> None:
    # Re-audit untouched production immediately before any write.
    plan = build_plan(PRODUCTION_WORKBOOK)
    print_plan(plan)

    if TEMP_WORKBOOK.exists():
        TEMP_WORKBOOK.unlink()

    # Perform the complete remediation and verification on a temp copy first.
    _, meta = _apply_to_copy(PRODUCTION_WORKBOOK, TEMP_WORKBOOK)

    backup = _production_backup_path()
    shutil.copy2(PRODUCTION_WORKBOOK, backup)

    # Atomic-ish same-directory replacement only after all verification passes.
    TEMP_WORKBOOK.replace(PRODUCTION_WORKBOOK)

    print()
    print("PASS: HISTORICAL OBSERVATION PRODUCTION REMEDIATION COMPLETED")
    print(f"Workbook       : {PRODUCTION_WORKBOOK}")
    print(f"Backup         : {backup}")
    print(f"Maintenance row: {meta['maintenance_row']}")
    print(f"Executed at    : {meta['execution_time'].isoformat(timespec='seconds')}")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Controlled remediation of Observation Workbook history for 2026-09-15"
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--audit",
        action="store_true",
        help="Read-only audit and recomputation plan (default)",
    )
    mode.add_argument(
        "--write-test",
        action="store_true",
        help="Apply remediation to a dedicated test copy only",
    )
    mode.add_argument(
        "--write-production",
        action="store_true",
        help="Apply remediation to production workbook",
    )
    parser.add_argument(
        "--confirm-production-write",
        action="store_true",
        help="Required second gate for --write-production",
    )
    args = parser.parse_args()

    if args.confirm_production_write and not args.write_production:
        parser.error(
            "--confirm-production-write is valid only with --write-production"
        )
    if args.write_production and not args.confirm_production_write:
        parser.error(
            "--write-production requires --confirm-production-write"
        )

    if args.write_test:
        plan, meta = _apply_to_copy(PRODUCTION_WORKBOOK, TEST_WORKBOOK)
        print_plan(plan)
        print()
        print("PASS: HISTORICAL OBSERVATION TEST REMEDIATION COMPLETED")
        print(f"Test workbook  : {TEST_WORKBOOK}")
        print(f"Maintenance row: {meta['maintenance_row']}")
        print(f"Executed at    : {meta['execution_time'].isoformat(timespec='seconds')}")
        print("PRODUCTION WORKBOOK WAS NOT MODIFIED")
        return 0

    if args.write_production:
        write_production()
        return 0

    plan = build_plan(PRODUCTION_WORKBOOK)
    print_plan(plan)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
