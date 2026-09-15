from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import historical_observation_remediation_2026_09_15 as subject
from observation_workbook import (
    DAILY_RUN_HEADERS,
    CANDIDATE_TRACKING_HEADERS,
    MAINTENANCE_LOG_HEADERS,
)


def make_fixture(path: Path) -> None:
    wb = Workbook()

    daily = wb.active
    daily.title = "Daily_Run"
    for col, header in enumerate(DAILY_RUN_HEADERS, start=1):
        daily.cell(row=4, column=col).value = header

    daily_values = {
        "Date": "2026-09-15",
        "Version": "v3.9.0",
        "PipelineStatus": "PASS",
        "PassSteps": "20/20",
        "RunId": "candidate-20260914-93aea4566066",
        "AsOfDate": "2026-09-14",
        "UniverseVersion": "AI_investing_universe_150_V2",
        "ScoreModelVersion": "technical-score-v3.8.1-r1",
        "RiskModelVersion": "risk-model-v3.8.2-p2",
        "UniverseConfigured": 150,
        "Ready": 144,
        "Excluded": 6,
        "ProviderRejected": 4,
        "StaleMarketData": 0,
        "InsufficientHistory": 2,
        "ExcludedSymbols": (
            "SKHY (INSUFFICIENT_HISTORY), "
            "SPCX (INSUFFICIENT_HISTORY), "
            "TSM (PROVIDER_REJECTED), "
            "LMND (PROVIDER_REJECTED), "
            "MCD (PROVIDER_REJECTED), "
            "DIS (PROVIDER_REJECTED)"
        ),
        "CoverageStatus": "PARTIAL",
        "BUY": 0,
        "WATCH": 2,
        "IGNORE": 142,
        "FinalStatus": "NO_ACTION",
        "Notes": "AUTOMATED Authority mapping; legacy fixed note",
    }
    daily_cols = {header: i + 1 for i, header in enumerate(DAILY_RUN_HEADERS)}
    for field, value in daily_values.items():
        daily.cell(row=5, column=daily_cols[field]).value = value

    candidates = wb.create_sheet("Candidate_Tracking")
    for col, header in enumerate(CANDIDATE_TRACKING_HEADERS, start=1):
        candidates.cell(row=4, column=col).value = header
    candidate_cols = {
        header: i + 1 for i, header in enumerate(CANDIDATE_TRACKING_HEADERS)
    }

    rows = [
        {
            "Date": "2026-09-15",
            "Ticker": "OKTA",
            "Signal": "WATCH",
            "Rank": 1,
            "FinalScore": 71.07639886675065,
            "FundamentalScore": None,
            "CombinedScore": None,
            "Price": 186.4499969482422,
            "WhyTracked": "OKTA technical reason",
            "ResearchNote": "OKTA approved research note",
            "30D": None,
            "60D": None,
            "90D": None,
            "Outcome": "OPEN",
        },
        {
            "Date": "2026-09-15",
            "Ticker": "CRWD",
            "Signal": "WATCH",
            "Rank": 2,
            "FinalScore": 66.68654844242202,
            "FundamentalScore": None,
            "CombinedScore": None,
            "Price": 235.3800048828125,
            "WhyTracked": "CRWD technical reason",
            "ResearchNote": "CRWD approved research note",
            "30D": None,
            "60D": None,
            "90D": None,
            "Outcome": "OPEN",
        },
    ]
    for excel_row, data in zip((5, 6), rows):
        for field, value in data.items():
            candidates.cell(
                row=excel_row,
                column=candidate_cols[field],
            ).value = value

    maintenance = wb.create_sheet("Maintenance_Log")
    for col, header in enumerate(MAINTENANCE_LOG_HEADERS, start=1):
        maintenance.cell(row=4, column=col).value = header

    # Reproduce production-style reserved/preformatted blank rows.
    for row in range(5, 9):
        maintenance.row_dimensions[row].height = 18

    # Extra unrelated sheet proves the whole-workbook protection check tolerates
    # unrelated historical content and keeps it unchanged.
    unrelated = wb.create_sheet("Unrelated_History")
    unrelated["A1"] = "DO NOT CHANGE"
    unrelated["B2"] = 12345

    wb.save(path)


class HistoricalObservationRemediationTests(unittest.TestCase):
    def test_scope_is_frozen(self):
        self.assertEqual(
            subject.DAILY_ALLOWED_FIELDS,
            {"RiskModelVersion", "CoverageStatus", "Notes"},
        )
        self.assertEqual(
            subject.CANDIDATE_ALLOWED_FIELDS,
            {"FundamentalScore", "CombinedScore"},
        )
        self.assertEqual(subject.TARGET_TICKERS, ("OKTA", "CRWD"))

    def test_target_identity_is_frozen(self):
        self.assertEqual(subject.TARGET_DATE.isoformat(), "2026-09-15")
        self.assertEqual(
            subject.TARGET_RUN_ID,
            "candidate-20260914-93aea4566066",
        )

    def test_remediation_commit_evidence_is_explicit(self):
        self.assertEqual(
            subject.REMEDIATION_COMMITS,
            {
                "RiskModelVersion": "821eff70",
                "Notes": "84d1f0b",
                "CoverageStatus": "30915ca",
                "MissingScores": "d757f31",
            },
        )

    def test_blank_helper(self):
        self.assertTrue(subject._is_blank(None))
        self.assertTrue(subject._is_blank(""))
        self.assertTrue(subject._is_blank("   "))
        self.assertFalse(subject._is_blank("MISSING"))
        self.assertFalse(subject._is_blank(0))

    def test_build_plan_recomputes_expected_values_from_existing_inputs(self):
        with tempfile.TemporaryDirectory() as td:
            workbook = Path(td) / "fixture.xlsx"
            make_fixture(workbook)

            plan = subject.build_plan(workbook)

            self.assertEqual(plan["daily_row"], 5)
            self.assertEqual(
                plan["daily_before"]["RiskModelVersion"],
                "risk-model-v3.8.2-p2",
            )
            self.assertEqual(
                plan["daily_after"]["RiskModelVersion"],
                "MISSING",
            )
            self.assertEqual(
                plan["daily_after"]["CoverageStatus"],
                "PARTIAL_PROVIDER_REJECTION",
            )
            self.assertEqual(
                plan["daily_after"]["Notes"],
                (
                    "Provider rejection: TSM, LMND, MCD, DIS (4); "
                    "Stale market data: none; "
                    "Insufficient history: SKHY, SPCX (2); "
                    "Readiness 144/150; excluded 6."
                ),
            )

            for ticker in ("OKTA", "CRWD"):
                _, before, after = plan["candidates"][ticker]
                self.assertIsNone(before["FundamentalScore"])
                self.assertIsNone(before["CombinedScore"])
                self.assertEqual(after["FundamentalScore"], "MISSING")
                self.assertEqual(after["CombinedScore"], "MISSING")

    def test_apply_to_copy_changes_only_allowed_cells_and_appends_maintenance(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "source.xlsx"
            destination = Path(td) / "destination.xlsx"
            make_fixture(source)

            before = load_workbook(source, data_only=False)
            try:
                before_daily = {
                    cell.coordinate: cell.value
                    for row in before["Daily_Run"].iter_rows()
                    for cell in row
                }
                before_candidates = {
                    cell.coordinate: cell.value
                    for row in before["Candidate_Tracking"].iter_rows()
                    for cell in row
                }
                before_unrelated = {
                    cell.coordinate: cell.value
                    for row in before["Unrelated_History"].iter_rows()
                    for cell in row
                }
            finally:
                before.close()

            plan, meta = subject._apply_to_copy(source, destination)

            after = load_workbook(destination, data_only=False)
            try:
                daily = after["Daily_Run"]
                daily_cols = {
                    daily.cell(row=4, column=col).value: col
                    for col in range(1, daily.max_column + 1)
                }
                self.assertEqual(
                    daily.cell(
                        row=plan["daily_row"],
                        column=daily_cols["RiskModelVersion"],
                    ).value,
                    "MISSING",
                )
                self.assertEqual(
                    daily.cell(
                        row=plan["daily_row"],
                        column=daily_cols["CoverageStatus"],
                    ).value,
                    "PARTIAL_PROVIDER_REJECTION",
                )

                candidates = after["Candidate_Tracking"]
                candidate_cols = {
                    candidates.cell(row=4, column=col).value: col
                    for col in range(1, candidates.max_column + 1)
                }
                for ticker in ("OKTA", "CRWD"):
                    row, _, _ = plan["candidates"][ticker]
                    self.assertEqual(
                        candidates.cell(
                            row=row,
                            column=candidate_cols["FundamentalScore"],
                        ).value,
                        "MISSING",
                    )
                    self.assertEqual(
                        candidates.cell(
                            row=row,
                            column=candidate_cols["CombinedScore"],
                        ).value,
                        "MISSING",
                    )

                # Explicitly prove protected semantic fields remain untouched.
                self.assertEqual(
                    candidates.cell(
                        row=5,
                        column=candidate_cols["WhyTracked"],
                    ).value,
                    "OKTA technical reason",
                )
                self.assertEqual(
                    candidates.cell(
                        row=5,
                        column=candidate_cols["ResearchNote"],
                    ).value,
                    "OKTA approved research note",
                )
                self.assertEqual(
                    candidates.cell(
                        row=6,
                        column=candidate_cols["WhyTracked"],
                    ).value,
                    "CRWD technical reason",
                )
                self.assertEqual(
                    candidates.cell(
                        row=6,
                        column=candidate_cols["ResearchNote"],
                    ).value,
                    "CRWD approved research note",
                )

                # Unrelated workbook history remains unchanged.
                after_unrelated = {
                    cell.coordinate: cell.value
                    for row in after["Unrelated_History"].iter_rows()
                    for cell in row
                }
                self.assertEqual(after_unrelated, before_unrelated)

                # Maintenance log was appended exactly once.
                maintenance = after["Maintenance_Log"]
                maintenance_cols = {
                    maintenance.cell(row=4, column=col).value: col
                    for col in range(1, maintenance.max_column + 1)
                }
                self.assertEqual(meta["maintenance_row"], 5)
                self.assertEqual(
                    maintenance.cell(
                        row=5,
                        column=maintenance_cols["Category"],
                    ).value,
                    "CODE_BUG",
                )
                self.assertEqual(
                    maintenance.cell(
                        row=5,
                        column=maintenance_cols["Action"],
                    ).value,
                    "CODE_FIX",
                )
                self.assertEqual(
                    maintenance.cell(
                        row=5,
                        column=maintenance_cols["Result"],
                    ).value,
                    "RESOLVED",
                )
                self.assertEqual(
                    maintenance.cell(
                        row=5,
                        column=maintenance_cols["Status"],
                    ).value,
                    "CLOSED",
                )

                # Compare all pre-existing Daily_Run cells, permitting only the
                # three explicitly allowed target cells to differ.
                allowed_daily_coords = {
                    daily.cell(
                        row=5, column=daily_cols["RiskModelVersion"]
                    ).coordinate,
                    daily.cell(
                        row=5, column=daily_cols["CoverageStatus"]
                    ).coordinate,
                    daily.cell(
                        row=5, column=daily_cols["Notes"]
                    ).coordinate,
                }
                after_daily = {
                    cell.coordinate: cell.value
                    for row in daily.iter_rows()
                    for cell in row
                }
                for coord, before_value in before_daily.items():
                    if coord in allowed_daily_coords:
                        continue
                    self.assertEqual(
                        after_daily.get(coord),
                        before_value,
                        f"Unexpected Daily_Run change at {coord}",
                    )

                # Compare all pre-existing Candidate_Tracking cells, permitting
                # only the four score cells to differ.
                allowed_candidate_coords = set()
                for ticker in ("OKTA", "CRWD"):
                    row, _, _ = plan["candidates"][ticker]
                    for field in ("FundamentalScore", "CombinedScore"):
                        allowed_candidate_coords.add(
                            candidates.cell(
                                row=row,
                                column=candidate_cols[field],
                            ).coordinate
                        )

                after_candidates = {
                    cell.coordinate: cell.value
                    for row in candidates.iter_rows()
                    for cell in row
                }
                for coord, before_value in before_candidates.items():
                    if coord in allowed_candidate_coords:
                        continue
                    self.assertEqual(
                        after_candidates.get(coord),
                        before_value,
                        f"Unexpected Candidate_Tracking change at {coord}",
                    )
            finally:
                after.close()


if __name__ == "__main__":
    unittest.main()
