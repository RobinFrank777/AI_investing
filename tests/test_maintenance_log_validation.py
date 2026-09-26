from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

import observation_workbook as m


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "_local_today", lambda: date(2026, 9, 26))
    path = tmp_path / "maintenance.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = m.MAINTENANCE_LOG_SHEET
    for col, header in enumerate(m.MAINTENANCE_LOG_HEADERS, 1):
        ws.cell(4, col, header)
    wb.save(path)
    wb.close()
    return path


def record(**changes):
    values = dict(date=date(2026, 9, 26), category="market_data", severity="p2",
                  ticker_scope=" all ", issue="Test issue", evidence="Test evidence",
                  status="closed")
    values.update(changes)
    return values


@pytest.mark.parametrize("changes,field", [
    ({"category": "DATA_QUALITY"}, "Category"),
    ({"severity": "P4"}, "Severity"),
    ({"status": "DONE"}, "Status"),
    *[({"action": value}, "Action") for value in
      ("", "   ", "DEFER_TO_BT3", "MONITOR_DISPLAY_LAYER", "update_data", 1)],
    *[({"result": value}, "Result") for value in
      ("", "   ", "PASSED", "OPEN", "PARTIAL_CLOSURE", "resolved", 1)],
    *[({"date": value}, "Date") for value in
      (None, "2026-09-26", 1, 1.5, object(), date(2026, 9, 27))],
    *[({"follow_up_date": value}, "FollowUpDate") for value in
      ("2026-09-27", 1, 1.5, object(), date(2026, 9, 25))],
])
def test_invalid_values_fail_before_open_and_preserve_bytes(workbook, monkeypatch, changes, field):
    before = workbook.read_bytes()
    def must_not_open(*args, **kwargs):
        pytest.fail("Validation must reject before opening workbook")
    monkeypatch.setattr(m, "load_workbook", must_not_open)
    with pytest.raises(RuntimeError, match=f"Maintenance_Log {field}"):
        m.append_maintenance_log(workbook, **record(**changes))
    assert workbook.read_bytes() == before


@pytest.mark.parametrize("action", [None, "UPDATE_DATA", "  UPDATE_DATA  ",
    "UPDATE_FUNDAMENTALS", "UPDATE_PROFILE", "RE-RUN_PIPELINE", "FIX_INPUT_FORMAT",
    "NO_CHANGE_OBSERVE", "DEFER_TO_NEXT_VERSION", "CODE_FIX"])
@pytest.mark.parametrize("result", [None, "RESOLVED", "  RESOLVED  ",
    "PARTIAL", "NO_CHANGE", "FAILED", "DEFERRED"])
def test_vocabulary_roundtrip(workbook, action, result):
    row = m.append_maintenance_log(workbook, **record(action=action, result=result))
    wb = load_workbook(workbook)
    ws = wb[m.MAINTENANCE_LOG_SHEET]
    assert row == 5
    assert [ws.cell(row, c).value for c in (2, 3, 4, 13)] == ["MARKET_DATA", "P2", "ALL", "CLOSED"]
    assert ws.cell(row, 7).value == (action.strip() if action is not None else None)
    assert ws.cell(row, 11).value == (result.strip() if result is not None else None)
    assert ws.cell(row, 12).value is None
    m.validate_maintenance_log_history_contiguous(ws)
    wb.close()


@pytest.mark.parametrize("day", [date(2026, 9, 26), datetime(2026, 9, 26, 14, 30)])
@pytest.mark.parametrize("follow", [None, date(2026, 9, 26), datetime(2026, 9, 26, 1),
                                   date(2027, 1, 1), datetime(2027, 1, 1, 15, 30)])
def test_dates_roundtrip_at_midnight(workbook, day, follow):
    row = m.append_maintenance_log(workbook, **record(date=day, follow_up_date=follow))
    wb = load_workbook(workbook)
    ws = wb[m.MAINTENANCE_LOG_SHEET]
    assert ws.cell(row, 1).value == datetime(2026, 9, 26)
    assert ws.cell(row, 12).value == (datetime.combine(follow, datetime.min.time()) if follow else None)
    wb.close()


def test_contiguous_append_duplicate_and_gap_fail_closed(workbook):
    m.append_maintenance_log(workbook, **record())
    before = workbook.read_bytes()
    with pytest.raises(RuntimeError):
        m.append_maintenance_log(workbook, **record())
    assert workbook.read_bytes() == before
    assert m.append_maintenance_log(workbook, **record(issue="Second issue")) == 6
    wb = load_workbook(workbook)
    ws = wb[m.MAINTENANCE_LOG_SHEET]
    m.validate_maintenance_log_history_contiguous(ws)
    ws.cell(8, 1, datetime(2026, 9, 26))
    wb.save(workbook)
    wb.close()
    before = workbook.read_bytes()
    with pytest.raises(RuntimeError):
        m.append_maintenance_log(workbook, **record(issue="Third issue"))
    assert workbook.read_bytes() == before


@pytest.mark.parametrize("field,value,message", [
    ("category", "DATA_QUALITY", "Maintenance_Log Category must be one of: " + ", ".join(sorted(m.MAINTENANCE_LOG_CATEGORIES))),
    ("severity", "P4", "Maintenance_Log Severity must be one of P0/P1/P2/P3."),
    ("status", "DONE", "Maintenance_Log Status must be one of OPEN/MONITORING/DEFERRED/CLOSED."),
])
def test_existing_error_messages_unchanged(field, value, message):
    with pytest.raises(RuntimeError) as exc:
        m.validate_maintenance_log_values(**record(**{field: value}))
    assert str(exc.value) == message
