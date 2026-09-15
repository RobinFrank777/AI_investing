"""Observation Workbook authority mapper and guarded workbook writer.

Phase 4:
- reads current production artifacts
- validates run identity consistency
- prints Daily_Run and Candidate_Tracking mappings
- --dry-run performs no workbook writes
- --write-test writes ONLY AI_investing_observation_test.xlsx
- --write-production requires --confirm-production-write
- production writes use a same-directory temp file, pre-write backup,
  duplicate gates, existing workbook writers, and atomic replacement
"""

from __future__ import annotations

import argparse
import math
import shutil
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import PROJECT_VERSION, REPO_ROOT
from current_run_status import load_current_run_status
from observation_workbook import append_candidate_tracking, append_daily_run


RESULTS_DIR = REPO_ROOT / "results"
CANDIDATES_PATH = RESULTS_DIR / "production_candidates.csv"
ACTION_REPORT_PATH = RESULTS_DIR / "portfolio_action_report.txt"
STOCK_RANK_PATH = RESULTS_DIR / "stock_rank.csv"
COMBINED_SCORE_PATH = RESULTS_DIR / "combined_score.csv"

OBSERVATION_DIR = Path.home() / "Documents" / "AI_investing_observation"
PRODUCTION_WORKBOOK_PATH = OBSERVATION_DIR / "AI_investing_observation.xlsx"
PRODUCTION_TEMP_PATH = (
    OBSERVATION_DIR / "AI_investing_observation.write-production.tmp.xlsx"
)
PRODUCTION_BACKUP_DIR = OBSERVATION_DIR / "backups"
TEST_WORKBOOK_PATH = OBSERVATION_DIR / "AI_investing_observation_test.xlsx"
TEST_TEMP_PATH = OBSERVATION_DIR / "AI_investing_observation_test.write-test.tmp.xlsx"

TRACKED_SIGNALS = {"BUY", "WATCH"}


def parse_colon_report(path: Path) -> dict[str, str]:
    if not path.is_file():
        raise FileNotFoundError(path)

    values: dict[str, str] = {}
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        if ":" not in raw_line:
            continue
        key, value = raw_line.split(":", 1)
        key = key.strip()
        value = value.strip()
        if key and key not in values:
            values[key] = value
    return values


def require_report_value(report: dict[str, str], key: str) -> str:
    value = str(report.get(key, "")).strip()
    if not value:
        raise RuntimeError(f"Missing required action-report field: {key}")
    return value


def require_unique(frame: pd.DataFrame, column: str) -> str:
    if column not in frame.columns:
        raise RuntimeError(f"production_candidates.csv missing column: {column}")

    values = frame[column].dropna().astype(str).str.strip()
    values = [value for value in values.unique() if value]

    if len(values) != 1:
        raise RuntimeError(
            f"Expected exactly one {column} in production_candidates.csv; "
            f"found {values}"
        )
    return values[0]


def parse_int(value: str, field_name: str) -> int:
    try:
        return int(str(value).strip())
    except ValueError as exc:
        raise RuntimeError(
            f"Expected integer for {field_name}; got {value!r}"
        ) from exc


def require_run_date() -> date:
    status = load_current_run_status()
    if not status:
        raise RuntimeError("current_run_status.json is missing or invalid")

    current_run_id = str(status.get("CurrentRunId", "")).strip()
    if not current_run_id:
        raise RuntimeError("CurrentRunId is missing from current_run_status.json")

    if PRODUCTION_WORKBOOK_PATH.is_file():
        history = pd.read_excel(
            PRODUCTION_WORKBOOK_PATH,
            sheet_name="Daily_Run",
            header=3,
            usecols=["Date", "RunId"],
        )

        matches = history.loc[
            history["RunId"]
            .fillna("")
            .astype(str)
            .str.strip()
            .eq(current_run_id),
            "Date",
        ].dropna()

        if not matches.empty:
            dates = (
                pd.to_datetime(matches, errors="coerce")
                .dropna()
                .dt.date
                .unique()
                .tolist()
            )

            if len(dates) != 1:
                raise RuntimeError(
                    f"RunId {current_run_id!r} has ambiguous Observation dates: "
                    f"{dates}"
                )

            return dates[0]

    start_time = str(status.get("StartTime", "")).strip()
    if not start_time:
        raise RuntimeError("StartTime is missing from current_run_status.json")

    try:
        return datetime.fromisoformat(start_time).date()
    except ValueError as exc:
        raise RuntimeError(f"Invalid StartTime: {start_time!r}") from exc


def build_daily_run_preview() -> dict[str, object]:
    status = load_current_run_status()
    if not status:
        raise RuntimeError("current_run_status.json is missing or invalid")

    run_date = require_run_date()

    if status.get("OverallRunStatus") != "PASS":
        raise RuntimeError(
            "Latest production pipeline is not PASS; refusing Daily_Run mapping"
        )

    passed_steps = status.get("PassedSteps")
    total_steps = status.get("TotalSteps")
    if not isinstance(passed_steps, int) or not isinstance(total_steps, int):
        raise RuntimeError(
            "PassedSteps / TotalSteps are missing from current_run_status.json"
        )
    if total_steps <= 0 or passed_steps < 0:
        raise RuntimeError(
            f"Invalid pipeline step counts: {passed_steps}/{total_steps}"
        )

    if passed_steps != total_steps:
        raise RuntimeError(
            f"Pipeline is not complete: {passed_steps}/{total_steps} PASS"
        )

    frame = pd.read_csv(CANDIDATES_PATH)
    if frame.empty:
        raise RuntimeError("production_candidates.csv is empty")

    candidate_run_id = require_unique(frame, "RunId")
    candidate_as_of = require_unique(frame, "AsOfDate")
    universe_version = require_unique(frame, "UniverseVersion")
    score_model_version = require_unique(frame, "ScoreModelVersion")

    status_run_id = str(status.get("CurrentRunId", "")).strip()
    status_as_of = str(status.get("AsOfDate", "")).strip()

    if status_run_id != candidate_run_id:
        raise RuntimeError(
            f"RunId mismatch: status={status_run_id!r}, "
            f"candidates={candidate_run_id!r}"
        )
    if status_as_of != candidate_as_of:
        raise RuntimeError(
            f"AsOfDate mismatch: status={status_as_of!r}, "
            f"candidates={candidate_as_of!r}"
        )

    report = parse_colon_report(ACTION_REPORT_PATH)

    report_run_id = require_report_value(report, "RunId")
    report_as_of = require_report_value(report, "AsOfDate")
    report_universe = require_report_value(report, "UniverseVersion")
    report_score_model = require_report_value(report, "ScoreModelVersion")

    identity_checks = {
        "RunId": (candidate_run_id, report_run_id),
        "AsOfDate": (candidate_as_of, report_as_of),
        "UniverseVersion": (universe_version, report_universe),
        "ScoreModelVersion": (score_model_version, report_score_model),
    }
    for field_name, (left, right) in identity_checks.items():
        if left != right:
            raise RuntimeError(
                f"{field_name} mismatch: candidates={left!r}, "
                f"action_report={right!r}"
            )

    if "TradeSignal" not in frame.columns:
        raise RuntimeError(
            "production_candidates.csv missing column: TradeSignal"
        )

    signal_counts = (
        frame["TradeSignal"]
        .dropna()
        .astype(str)
        .str.strip()
        .str.upper()
        .value_counts()
    )
    buy = int(signal_counts.get("BUY", 0))
    watch = int(signal_counts.get("WATCH", 0))
    ignore = int(signal_counts.get("IGNORE", 0))

    configured = parse_int(
        require_report_value(report, "Configured Universe"),
        "Configured Universe",
    )
    ready = parse_int(
        require_report_value(report, "Research Ready"),
        "Research Ready",
    )
    excluded = parse_int(
        require_report_value(report, "Excluded"),
        "Excluded",
    )
    provider_rejected = parse_int(
        require_report_value(report, "Provider Rejected"),
        "Provider Rejected",
    )
    stale_market_data = parse_int(
        require_report_value(report, "Stale Market Data"),
        "Stale Market Data",
    )
    insufficient_history = parse_int(
        require_report_value(report, "Insufficient History"),
        "Insufficient History",
    )

    if ready + excluded != configured:
        raise RuntimeError(
            f"Coverage arithmetic mismatch: {ready} + {excluded} != {configured}"
        )

    if buy + watch + ignore != ready:
        raise RuntimeError(
            "Signal arithmetic mismatch: "
            f"{buy} + {watch} + {ignore} != {ready}"
        )

    final_status = require_report_value(report, "Portfolio Decision")
    report_status = require_report_value(report, "Report Status")
    if final_status != report_status:
        raise RuntimeError(
            f"Action report status mismatch: "
            f"Portfolio Decision={final_status!r}, Report Status={report_status!r}"
        )

    return {
        "Date": run_date,
        "Version": PROJECT_VERSION,
        "PipelineStatus": status["OverallRunStatus"],
        "PassSteps": f"{passed_steps}/{total_steps}",
        "RunId": candidate_run_id,
        "AsOfDate": date.fromisoformat(candidate_as_of),
        "UniverseVersion": universe_version,
        "ScoreModelVersion": score_model_version,
        "RiskModelVersion": require_report_value(report, "RiskModelVersion"),
        "UniverseConfigured": configured,
        "Ready": ready,
        "Excluded": excluded,
        "ProviderRejected": provider_rejected,
        "StaleMarketData": stale_market_data,
        "InsufficientHistory": insufficient_history,
        "ExcludedSymbols": require_report_value(report, "Excluded Symbols"),
        "CoverageStatus": require_report_value(report, "Data Coverage"),
        "BUY": buy,
        "WATCH": watch,
        "IGNORE": ignore,
        "FinalStatus": final_status,
        "Notes": (
            "Authority mapping; current_run_status.json + "
            "production_candidates.csv + portfolio_action_report.txt"
        ),
    }


def _require_columns(frame: pd.DataFrame, columns: tuple[str, ...], label: str) -> None:
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise RuntimeError(
            f"{label} missing required columns: {', '.join(missing)}"
        )


def _optional_combined_values(
    ticker: str,
    *,
    run_id: str,
    as_of_date: str,
    universe_version: str,
    score_model_version: str,
) -> tuple[object, object]:
    if not COMBINED_SCORE_PATH.is_file():
        return None, None

    frame = pd.read_csv(COMBINED_SCORE_PATH)
    if frame.empty:
        return None, None

    _require_columns(
        frame,
        (
            "Ticker",
            "RunId",
            "AsOfDate",
            "UniverseVersion",
            "ScoreModelVersion",
            "FundamentalScore",
            "CombinedScore",
        ),
        "combined_score.csv",
    )

    matches = frame.loc[
        frame["Ticker"].astype(str).str.strip().str.upper().eq(ticker)
    ].copy()

    if matches.empty:
        return None, None
    if len(matches) != 1:
        raise RuntimeError(
            f"combined_score.csv expected one row for {ticker}; found {len(matches)}"
        )

    row = matches.iloc[0]
    checks = {
        "RunId": run_id,
        "AsOfDate": as_of_date,
        "UniverseVersion": universe_version,
        "ScoreModelVersion": score_model_version,
    }
    for field_name, expected in checks.items():
        actual = str(row[field_name]).strip()
        if actual != expected:
            raise RuntimeError(
                f"combined_score.csv {ticker} {field_name} mismatch: "
                f"{actual!r} != {expected!r}"
            )

    fundamental = row["FundamentalScore"]
    combined = row["CombinedScore"]
    fundamental = None if pd.isna(fundamental) else float(fundamental)
    combined = None if pd.isna(combined) else float(combined)
    return fundamental, combined


def build_candidate_tracking_previews() -> list[dict[str, object]]:
    candidates = pd.read_csv(CANDIDATES_PATH)
    run_date = require_run_date()
    if candidates.empty:
        raise RuntimeError("production_candidates.csv is empty")

    _require_columns(
        candidates,
        (
            "Ticker",
            "RunId",
            "AsOfDate",
            "CandidateRank",
            "FinalScore",
            "TradeSignal",
            "ScoreModelVersion",
            "UniverseVersion",
        ),
        "production_candidates.csv",
    )

    run_id = require_unique(candidates, "RunId")
    as_of_date = require_unique(candidates, "AsOfDate")
    universe_version = require_unique(candidates, "UniverseVersion")
    score_model_version = require_unique(candidates, "ScoreModelVersion")

    signals = candidates["TradeSignal"].fillna("").astype(str).str.strip().str.upper()
    tracked = candidates.loc[signals.isin(TRACKED_SIGNALS)].copy()
    tracked["_Signal"] = signals.loc[tracked.index]

    if tracked.empty:
        return []

    stock_rank = pd.read_csv(STOCK_RANK_PATH)
    if stock_rank.empty:
        raise RuntimeError("stock_rank.csv is empty")

    _require_columns(
        stock_rank,
        (
            "Ticker",
            "MarketDataDate",
            "Close",
            "FinalScore",
            "TradeSignal",
            "Reason",
            "ScoreModelVersion",
            "UniverseVersion",
        ),
        "stock_rank.csv",
    )

    previews: list[dict[str, object]] = []

    for _, candidate in tracked.sort_values(
        ["CandidateRank", "Ticker"], kind="mergesort"
    ).iterrows():
        ticker = str(candidate["Ticker"]).strip().upper()
        signal = str(candidate["_Signal"]).strip().upper()

        rank_value = candidate["CandidateRank"]
        if pd.isna(rank_value):
            raise RuntimeError(f"{ticker} CandidateRank is missing")
        rank = int(rank_value)
        if float(rank_value) != rank or rank < 1:
            raise RuntimeError(f"{ticker} CandidateRank is invalid: {rank_value!r}")

        final_score = float(candidate["FinalScore"])
        if not math.isfinite(final_score):
            raise RuntimeError(f"{ticker} FinalScore is not finite")

        matches = stock_rank.loc[
            stock_rank["Ticker"].astype(str).str.strip().str.upper().eq(ticker)
        ].copy()
        if len(matches) != 1:
            raise RuntimeError(
                f"stock_rank.csv expected exactly one row for {ticker}; "
                f"found {len(matches)}"
            )

        rank_row = matches.iloc[0]
        checks = {
            "MarketDataDate": as_of_date,
            "ScoreModelVersion": score_model_version,
            "UniverseVersion": universe_version,
            "TradeSignal": signal,
        }
        for field_name, expected in checks.items():
            actual = str(rank_row[field_name]).strip()
            if field_name == "TradeSignal":
                actual = actual.upper()
            if actual != expected:
                raise RuntimeError(
                    f"stock_rank.csv {ticker} {field_name} mismatch: "
                    f"{actual!r} != {expected!r}"
                )

        rank_final_score = float(rank_row["FinalScore"])
        if not math.isclose(
            final_score, rank_final_score, rel_tol=0.0, abs_tol=1e-9
        ):
            raise RuntimeError(
                f"stock_rank.csv {ticker} FinalScore mismatch: "
                f"{rank_final_score} != {final_score}"
            )

        price = float(rank_row["Close"])
        if not math.isfinite(price) or price <= 0:
            raise RuntimeError(f"stock_rank.csv {ticker} Close is invalid")

        why_tracked = str(rank_row["Reason"]).strip()
        if not why_tracked:
            raise RuntimeError(f"stock_rank.csv {ticker} Reason is missing")

        fundamental_score, combined_score = _optional_combined_values(
            ticker,
            run_id=run_id,
            as_of_date=as_of_date,
            universe_version=universe_version,
            score_model_version=score_model_version,
        )

        previews.append(
            {
                "Date": run_date,
                "Ticker": ticker,
                "Signal": signal,
                "Rank": rank,
                "FinalScore": final_score,
                "FundamentalScore": fundamental_score,
                "CombinedScore": combined_score,
                "Price": price,
                "WhyTracked": why_tracked,
                "ResearchNote": None,
                "30D": None,
                "60D": None,
                "90D": None,
                "Outcome": "OPEN",
            }
        )

    return previews


def print_previews(
    daily_preview: dict[str, object],
    candidate_previews: list[dict[str, object]],
) -> None:
    print("\\nOBSERVATION RUNNER — DAILY_RUN")
    print("=" * 72)
    for key, value in daily_preview.items():
        print(f"{key:<20}: {value}")
    print("=" * 72)

    print("\\nOBSERVATION RUNNER — CANDIDATE_TRACKING")
    print("=" * 72)
    if not candidate_previews:
        print("No BUY/WATCH candidates for this run.")
    else:
        for index, preview in enumerate(candidate_previews, start=1):
            if index > 1:
                print("-" * 72)
            for key, value in preview.items():
                print(f"{key:<20}: {value}")
    print("=" * 72)



def prepare_test_fixture(
    daily_preview: dict[str, object],
    candidate_previews: list[dict[str, object]],
) -> None:
    """Build a disposable fixture from the production workbook.

    The production workbook is read/copy source only. The current Daily_Run
    RunId and current Candidate_Tracking Date+Ticker rows are removed only
    from the temporary fixture so the existing append writers can be tested
    without disabling duplicate protection.
    """

    if PRODUCTION_WORKBOOK_PATH.name != "AI_investing_observation.xlsx":
        raise RuntimeError("Refusing fixture build: unexpected production workbook filename")
    if TEST_WORKBOOK_PATH.name != "AI_investing_observation_test.xlsx":
        raise RuntimeError("Refusing fixture build: unexpected test workbook filename")
    if PRODUCTION_WORKBOOK_PATH.resolve() == TEST_WORKBOOK_PATH.resolve():
        raise RuntimeError("Production and test workbook paths must be different")
    if not PRODUCTION_WORKBOOK_PATH.is_file():
        raise FileNotFoundError(PRODUCTION_WORKBOOK_PATH)

    TEST_WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)

    if TEST_TEMP_PATH.exists():
        TEST_TEMP_PATH.unlink()

    shutil.copy2(PRODUCTION_WORKBOOK_PATH, TEST_TEMP_PATH)

    wb = None
    try:
        wb = load_workbook(TEST_TEMP_PATH)

        if "Daily_Run" not in wb.sheetnames:
            raise RuntimeError("Test fixture missing Daily_Run sheet")
        if "Candidate_Tracking" not in wb.sheetnames:
            raise RuntimeError("Test fixture missing Candidate_Tracking sheet")

        daily_ws = wb["Daily_Run"]
        candidate_ws = wb["Candidate_Tracking"]

        target_run_id = str(daily_preview["RunId"]).strip()
        if not target_run_id:
            raise RuntimeError("Daily_Run RunId must not be empty")

        daily_rows = []
        for row in range(5, daily_ws.max_row + 1):
            existing_run_id = daily_ws.cell(row=row, column=5).value
            if str(existing_run_id or "").strip() == target_run_id:
                daily_rows.append(row)

        if len(daily_rows) > 1:
            raise RuntimeError(
                f"Production fixture contains duplicate Daily_Run RunId "
                f"{target_run_id!r}: rows {daily_rows}"
            )

        for row in reversed(daily_rows):
            daily_ws.delete_rows(row, 1)

        target_keys = {
            (item["Date"], str(item["Ticker"]).strip().upper())
            for item in candidate_previews
        }

        key_rows: dict[tuple[date, str], list[int]] = {
            key: [] for key in target_keys
        }

        for row in range(5, candidate_ws.max_row + 1):
            existing_date = candidate_ws.cell(row=row, column=1).value
            existing_ticker = candidate_ws.cell(row=row, column=2).value

            if existing_date is None or existing_ticker is None:
                continue

            try:
                existing_day = pd.to_datetime(existing_date).date()
            except (TypeError, ValueError):
                continue

            key = (
                existing_day,
                str(existing_ticker).strip().upper(),
            )
            if key in key_rows:
                key_rows[key].append(row)

        duplicate_keys = {
            key: rows for key, rows in key_rows.items() if len(rows) > 1
        }
        if duplicate_keys:
            raise RuntimeError(
                f"Production fixture contains duplicate Candidate_Tracking "
                f"Date+Ticker rows: {duplicate_keys}"
            )

        candidate_rows = sorted(
            (rows[0] for rows in key_rows.values() if rows),
            reverse=True,
        )
        for row in candidate_rows:
            candidate_ws.delete_rows(row, 1)

        wb.save(TEST_TEMP_PATH)

    except Exception:
        if wb is not None:
            wb.close()
        if TEST_TEMP_PATH.exists():
            TEST_TEMP_PATH.unlink()
        raise
    else:
        wb.close()


def _excel_roundtrip_datetime(value):
    """Use datetime values for Excel date fields so reopen verification is stable."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    return value

def write_test_workbook(
    daily_preview: dict[str, object],
    candidate_previews: list[dict[str, object]],
) -> tuple[int, list[int]]:
    if TEST_WORKBOOK_PATH.name != "AI_investing_observation_test.xlsx":
        raise RuntimeError("Refusing write: unexpected test workbook filename")

    prepare_test_fixture(
        daily_preview,
        candidate_previews,
    )

    try:
        daily_row = append_daily_run(
            file_path=TEST_TEMP_PATH,
            date=_excel_roundtrip_datetime(daily_preview["Date"]),
            version=daily_preview["Version"],
            pipeline_status=daily_preview["PipelineStatus"],
            pass_steps=daily_preview["PassSteps"],
            run_id=daily_preview["RunId"],
            as_of_date=_excel_roundtrip_datetime(daily_preview["AsOfDate"]),
            universe_version=daily_preview["UniverseVersion"],
            score_model_version=daily_preview["ScoreModelVersion"],
            risk_model_version=daily_preview["RiskModelVersion"],
            universe_configured=daily_preview["UniverseConfigured"],
            ready=daily_preview["Ready"],
            excluded=daily_preview["Excluded"],
            provider_rejected=daily_preview["ProviderRejected"],
            stale_market_data=daily_preview["StaleMarketData"],
            insufficient_history=daily_preview["InsufficientHistory"],
            excluded_symbols=daily_preview["ExcludedSymbols"],
            coverage_status=daily_preview["CoverageStatus"],
            buy=daily_preview["BUY"],
            watch=daily_preview["WATCH"],
            ignore=daily_preview["IGNORE"],
            final_status=daily_preview["FinalStatus"],
            notes="WRITE_TEST " + str(daily_preview["Notes"]),
        )

        candidate_rows = []
        for item in candidate_previews:
            candidate_rows.append(
                append_candidate_tracking(
                    file_path=TEST_TEMP_PATH,
                    date=_excel_roundtrip_datetime(item["Date"]),
                    ticker=item["Ticker"],
                    signal=item["Signal"],
                    rank=item["Rank"],
                    final_score=item["FinalScore"],
                    fundamental_score=item["FundamentalScore"],
                    combined_score=item["CombinedScore"],
                    price=item["Price"],
                    why_tracked=item["WhyTracked"],
                    research_note=item["ResearchNote"],
                    day_30=item["30D"],
                    day_60=item["60D"],
                    day_90=item["90D"],
                    outcome=item["Outcome"],
                )
            )

        TEST_TEMP_PATH.replace(TEST_WORKBOOK_PATH)
        return daily_row, candidate_rows

    except Exception:
        if TEST_TEMP_PATH.exists():
            TEST_TEMP_PATH.unlink()
        raise



def validate_production_write_gate(
    daily_preview: dict[str, object],
    candidate_previews: list[dict[str, object]],
) -> None:
    """Fail closed unless the production workbook is safe for one new append."""

    if PRODUCTION_WORKBOOK_PATH.name != "AI_investing_observation.xlsx":
        raise RuntimeError("Refusing production write: unexpected workbook filename")
    if not PRODUCTION_WORKBOOK_PATH.is_file():
        raise FileNotFoundError(PRODUCTION_WORKBOOK_PATH)
    if PRODUCTION_TEMP_PATH.resolve() == PRODUCTION_WORKBOOK_PATH.resolve():
        raise RuntimeError("Production temp path must differ from production workbook")

    wb = load_workbook(PRODUCTION_WORKBOOK_PATH, data_only=False, read_only=False)
    try:
        for required_sheet in ("Daily_Run", "Candidate_Tracking"):
            if required_sheet not in wb.sheetnames:
                raise RuntimeError(
                    f"Production workbook missing required sheet: {required_sheet}"
                )

        daily_ws = wb["Daily_Run"]
        candidate_ws = wb["Candidate_Tracking"]

        target_date = daily_preview["Date"]
        target_run_id = str(daily_preview["RunId"]).strip()

        if not target_run_id:
            raise RuntimeError("Production write requires non-empty RunId")

        daily_date_matches = []
        daily_run_id_matches = []

        for row in range(5, daily_ws.max_row + 1):
            existing_date = daily_ws.cell(row=row, column=1).value
            existing_run_id = daily_ws.cell(row=row, column=5).value

            if existing_date is not None:
                try:
                    existing_day = pd.to_datetime(existing_date).date()
                except (TypeError, ValueError):
                    existing_day = None
                if existing_day == target_date:
                    daily_date_matches.append(row)

            if str(existing_run_id or "").strip() == target_run_id:
                daily_run_id_matches.append(row)

        if daily_date_matches:
            raise RuntimeError(
                f"Production write blocked: Daily_Run Date {target_date} "
                f"already exists at rows {daily_date_matches}"
            )

        if daily_run_id_matches:
            raise RuntimeError(
                f"Production write blocked: RunId {target_run_id!r} "
                f"already exists at rows {daily_run_id_matches}"
            )

        candidate_targets = {
            (item["Date"], str(item["Ticker"]).strip().upper())
            for item in candidate_previews
        }

        candidate_matches: dict[tuple[date, str], list[int]] = {
            key: [] for key in candidate_targets
        }

        for row in range(5, candidate_ws.max_row + 1):
            existing_date = candidate_ws.cell(row=row, column=1).value
            existing_ticker = candidate_ws.cell(row=row, column=2).value

            if existing_date is None or existing_ticker is None:
                continue

            try:
                existing_day = pd.to_datetime(existing_date).date()
            except (TypeError, ValueError):
                continue

            key = (
                existing_day,
                str(existing_ticker).strip().upper(),
            )
            if key in candidate_matches:
                candidate_matches[key].append(row)

        existing_candidate_keys = {
            key: rows for key, rows in candidate_matches.items() if rows
        }
        if existing_candidate_keys:
            raise RuntimeError(
                "Production write blocked: Candidate_Tracking Date+Ticker "
                f"already exists: {existing_candidate_keys}"
            )

    finally:
        wb.close()


def _make_production_backup_path() -> Path:
    """Create a unique timestamped backup path without modifying the workbook."""
    timestamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    path = (
        PRODUCTION_BACKUP_DIR
        / f"AI_investing_observation.before-production-write-{timestamp}.xlsx"
    )
    if path.exists():
        raise RuntimeError(f"Backup path already exists: {path}")
    return path


def write_production_workbook(
    daily_preview: dict[str, object],
    candidate_previews: list[dict[str, object]],
) -> tuple[int, list[int], Path]:
    """Transactionally append the current authority mapping to production."""

    validate_production_write_gate(daily_preview, candidate_previews)

    PRODUCTION_BACKUP_DIR.mkdir(parents=True, exist_ok=True)

    if PRODUCTION_TEMP_PATH.exists():
        PRODUCTION_TEMP_PATH.unlink()

    shutil.copy2(PRODUCTION_WORKBOOK_PATH, PRODUCTION_TEMP_PATH)

    backup_path = _make_production_backup_path()

    try:
        daily_row = append_daily_run(
            file_path=PRODUCTION_TEMP_PATH,
            date=_excel_roundtrip_datetime(daily_preview["Date"]),
            version=daily_preview["Version"],
            pipeline_status=daily_preview["PipelineStatus"],
            pass_steps=daily_preview["PassSteps"],
            run_id=daily_preview["RunId"],
            as_of_date=_excel_roundtrip_datetime(daily_preview["AsOfDate"]),
            universe_version=daily_preview["UniverseVersion"],
            score_model_version=daily_preview["ScoreModelVersion"],
            risk_model_version=daily_preview["RiskModelVersion"],
            universe_configured=daily_preview["UniverseConfigured"],
            ready=daily_preview["Ready"],
            excluded=daily_preview["Excluded"],
            provider_rejected=daily_preview["ProviderRejected"],
            stale_market_data=daily_preview["StaleMarketData"],
            insufficient_history=daily_preview["InsufficientHistory"],
            excluded_symbols=daily_preview["ExcludedSymbols"],
            coverage_status=daily_preview["CoverageStatus"],
            buy=daily_preview["BUY"],
            watch=daily_preview["WATCH"],
            ignore=daily_preview["IGNORE"],
            final_status=daily_preview["FinalStatus"],
            notes="AUTOMATED " + str(daily_preview["Notes"]),
        )

        candidate_rows = []
        for item in candidate_previews:
            candidate_rows.append(
                append_candidate_tracking(
                    file_path=PRODUCTION_TEMP_PATH,
                    date=_excel_roundtrip_datetime(item["Date"]),
                    ticker=item["Ticker"],
                    signal=item["Signal"],
                    rank=item["Rank"],
                    final_score=item["FinalScore"],
                    fundamental_score=item["FundamentalScore"],
                    combined_score=item["CombinedScore"],
                    price=item["Price"],
                    why_tracked=item["WhyTracked"],
                    research_note=item["ResearchNote"],
                    day_30=item["30D"],
                    day_60=item["60D"],
                    day_90=item["90D"],
                    outcome=item["Outcome"],
                )
            )

        # Backup the untouched production workbook only after every temp-file
        # append and post-save verification has passed.
        shutil.copy2(PRODUCTION_WORKBOOK_PATH, backup_path)

        # Same-directory replacement is atomic on the normal local filesystem.
        PRODUCTION_TEMP_PATH.replace(PRODUCTION_WORKBOOK_PATH)

        return daily_row, candidate_rows, backup_path

    except Exception:
        if PRODUCTION_TEMP_PATH.exists():
            PRODUCTION_TEMP_PATH.unlink()
        raise


def build_research_status(
    candidate_previews: list[dict[str, object]],
) -> list[dict[str, object]]:
    # Read-only ResearchNote completeness check for current BUY/WATCH candidates.

    if not PRODUCTION_WORKBOOK_PATH.is_file():
        raise FileNotFoundError(PRODUCTION_WORKBOOK_PATH)

    wb = load_workbook(PRODUCTION_WORKBOOK_PATH, data_only=False, read_only=False)
    try:
        if "Candidate_Tracking" not in wb.sheetnames:
            raise RuntimeError("Production workbook missing Candidate_Tracking sheet")

        ws = wb["Candidate_Tracking"]
        results: list[dict[str, object]] = []

        for item in candidate_previews:
            target_date = item["Date"]
            target_ticker = str(item["Ticker"]).strip().upper()
            matches: list[int] = []

            for row in range(5, ws.max_row + 1):
                existing_date = ws.cell(row=row, column=1).value
                existing_ticker = ws.cell(row=row, column=2).value

                if existing_date is None or existing_ticker is None:
                    continue

                try:
                    existing_day = pd.to_datetime(existing_date).date()
                except (TypeError, ValueError):
                    continue

                if (
                    existing_day == target_date
                    and str(existing_ticker).strip().upper() == target_ticker
                ):
                    matches.append(row)

            if len(matches) > 1:
                raise RuntimeError(
                    "Candidate_Tracking key is not unique for "
                    f"{target_date} / {target_ticker}: rows {matches}"
                )

            if not matches:
                status = "NOT_WRITTEN"
                row = None
            else:
                row = matches[0]
                note = ws.cell(row=row, column=10).value
                status = (
                    "COMPLETE"
                    if note is not None and str(note).strip()
                    else "RESEARCH_REQUIRED"
                )

            results.append(
                {
                    "Date": target_date,
                    "Ticker": target_ticker,
                    "Row": row,
                    "ResearchStatus": status,
                }
            )

        return results
    finally:
        wb.close()


def print_research_status(
    candidate_previews: list[dict[str, object]],
) -> list[dict[str, object]]:
    results = build_research_status(candidate_previews)

    print("\nOBSERVATION RUNNER — RESEARCH STATUS")
    print("=" * 72)

    if not results:
        print("No BUY/WATCH candidates for this run.")
    else:
        for item in results:
            print(
                f"{item['Date']}  {item['Ticker']:<8}  "
                f"{item['ResearchStatus']}"
            )

    required = [
        item for item in results
        if item["ResearchStatus"] == "RESEARCH_REQUIRED"
    ]
    not_written = [
        item for item in results
        if item["ResearchStatus"] == "NOT_WRITTEN"
    ]

    if required:
        print("-" * 72)
        print("RESEARCH_REQUIRED:")
        for item in required:
            print(f"  {item['Ticker']}")
    elif not_written:
        print("-" * 72)
        print("Research status is pending Candidate_Tracking production write.")
    else:
        print("-" * 72)
        print("RESEARCH COMPLETE")

    print("=" * 72)
    return results



def main() -> int:
    parser = argparse.ArgumentParser(
        description="Observation Workbook authority mapper / guarded writer"
    )
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate authorities and print mappings without workbook writes",
    )
    mode.add_argument(
        "--write-test",
        action="store_true",
        help="Write mappings to the hard-coded test workbook only",
    )
    mode.add_argument(
        "--write-production",
        action="store_true",
        help="Request a guarded production workbook write",
    )
    mode.add_argument(
        "--research-status",
        action="store_true",
        help="Read-only ResearchNote completeness check for current BUY/WATCH candidates",
    )
    parser.add_argument(
        "--confirm-production-write",
        action="store_true",
        help="Required second gate for --write-production",
    )
    args = parser.parse_args()

    if args.confirm_production_write and not args.write_production:
        parser.error(
            "--confirm-production-write is valid only with --write-production"
        )

    if args.write_production and not args.confirm_production_write:
        parser.error(
            "--write-production requires --confirm-production-write"
        )

    daily_preview = build_daily_run_preview()
    candidate_previews = build_candidate_tracking_previews()
    print_previews(daily_preview, candidate_previews)

    if args.research_status:
        print_research_status(candidate_previews)
        return 0

    if args.dry_run:
        print("PASS: authority checks completed")
        print("NO WORKBOOK WRITE WAS PERFORMED")
        return 0

    if args.write_test:
        daily_row, candidate_rows = write_test_workbook(
            daily_preview,
            candidate_previews,
        )

        print("\\nPASS: TEST WORKBOOK WRITE COMPLETED")
        print(f"Workbook             : {TEST_WORKBOOK_PATH}")
        print(f"Daily_Run row        : {daily_row}")
        print(f"Candidate rows       : {candidate_rows}")
        print("PRODUCTION WORKBOOK  : NOT TOUCHED")
        return 0

    daily_row, candidate_rows, backup_path = write_production_workbook(
        daily_preview,
        candidate_previews,
    )

    print("\\nPASS: PRODUCTION WORKBOOK WRITE COMPLETED")
    print(f"Workbook             : {PRODUCTION_WORKBOOK_PATH}")
    print(f"Backup               : {backup_path}")
    print(f"Daily_Run row        : {daily_row}")
    print(f"Candidate rows       : {candidate_rows}")
    print_research_status(candidate_previews)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
